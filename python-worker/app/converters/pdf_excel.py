"""Conversión PDF → Excel (.xlsx): extracción de TABLAS con PyMuPDF + openpyxl.

Un PDF no es una hoja de cálculo: esta operación detecta las tablas del
documento (find_tables de PyMuPDF, ya empaquetado para pdf_to_word) y las
vuelca a un libro de Excel, una hoja por tabla, con la página de origen en el
nombre de la hoja. El texto suelto fuera de tablas NO se exporta — la UI lo
avisa. Si el documento no tiene tablas detectables, error claro en vez de un
xlsx vacío.
"""
from pathlib import Path

import fitz  # PyMuPDF
from openpyxl import Workbook

from app.converters.pdf_word import assert_pdf_extractable

# Tope defensivo: un PDF patológico con cientos de "tablas" detectadas no debe
# producir un libro inmanejable ni comerse la RAM del worker.
MAX_TABLES = 200


def convert_pdf_to_excel(input_path: Path, output_path: Path) -> int:
    """Extrae las tablas del PDF a un .xlsx (una hoja por tabla).

    Devuelve cuántas tablas se exportaron. ValueError con mensaje claro si el
    PDF está cifrado/escaneado o no tiene tablas detectables.
    """
    assert_pdf_extractable(input_path)

    workbook = Workbook()
    workbook.remove(workbook.active)  # sin hoja vacía inicial
    exported = 0

    doc = fitz.open(str(input_path))
    try:
        for page_index, page in enumerate(doc, start=1):
            for table in page.find_tables().tables:
                if exported >= MAX_TABLES:
                    break
                rows = table.extract()
                if not _has_content(rows):
                    continue
                exported += 1
                sheet = workbook.create_sheet(_sheet_name(page_index, exported))
                for row in rows:
                    # Celdas None (huecos de celdas combinadas) → vacías.
                    sheet.append(['' if cell is None else cell for cell in row])
    finally:
        doc.close()

    if exported == 0:
        raise ValueError(
            'No tables were detected in the PDF; this tool extracts tables '
            '(for the full text use PDF to Word)')

    workbook.save(output_path)
    return exported


def _has_content(rows) -> bool:
    return any(cell not in (None, '') for row in rows for cell in row)


def _sheet_name(page_index: int, table_number: int) -> str:
    # Excel limita el nombre de hoja a 31 caracteres; este formato cabe siempre.
    return f'Tabla {table_number} (pag {page_index})'
