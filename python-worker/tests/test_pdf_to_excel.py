"""Tests de la operación pdf_to_excel (tablas de un PDF → xlsx).

Los PDFs con tabla se generan con reportlab/platypus (rejilla real que
find_tables detecta); la validación de cifrado/escaneado compartida con
pdf_to_word ya se prueba en test_pdf_to_word.py.
"""
from pathlib import Path

import pytest
from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import SimpleDocTemplate, Spacer, Table, TableStyle

from app.config import settings
from app.converters.pdf_excel import convert_pdf_to_excel
from app.operations.pdf_to_excel import handle_pdf_to_excel


class FakeCursor:
    def __init__(self, originals):
        self.originals = originals
        self.inserts = []

    def execute(self, sql, params=None):
        if sql.strip().upper().startswith('INSERT'):
            self.inserts.append(params)

    def fetchall(self):
        return list(self.originals)

    def fetchone(self):
        return self.originals[0] if self.originals else None

    def output_paths(self):
        return [Path(p[4]) for p in self.inserts]


@pytest.fixture(autouse=True)
def _output_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, 'OUTPUT_DIR', str(tmp_path / 'out'))
    yield


def _original(path: Path, name: str = 'doc.pdf'):
    return {'id': 'f1', 'file_path': str(path), 'original_filename': name,
            'filename': path.name}


def make_table_pdf(path: Path, tables_data) -> None:
    """PDF con una o más tablas con rejilla (una tras otra)."""
    doc = SimpleDocTemplate(str(path), pagesize=letter)
    flow = []
    for data in tables_data:
        table = Table(data)
        table.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 0.5, colors.black)]))
        flow.append(table)
        # Separación generosa: dos tablas pegadas se detectan como UNA sola.
        flow.append(Spacer(1, 60))
    doc.build(flow)


def make_text_pdf(path: Path) -> None:
    """PDF con texto pero SIN tablas."""
    c = rl_canvas.Canvas(str(path), pagesize=letter)
    c.setFont('Helvetica', 12)
    c.drawString(72, 700, 'Parrafo normal sin ninguna tabla.')
    c.save()


SOCIOS = [['Socio', 'Balance'], ['José Núñez', '87,500.50'], ['María Peña', '12,300.75']]


def test_extracts_table_with_accents(tmp_path):
    src = tmp_path / 'in.pdf'
    make_table_pdf(src, [SOCIOS])
    out = tmp_path / 'out.xlsx'
    tables = convert_pdf_to_excel(src, out)

    assert tables == 1
    wb = load_workbook(out)
    sheet = wb[wb.sheetnames[0]]
    assert sheet['A1'].value == 'Socio'
    assert sheet['A2'].value == 'José Núñez'
    assert sheet['B3'].value == '12,300.75'


def test_multiple_tables_one_sheet_each(tmp_path):
    src = tmp_path / 'in.pdf'
    make_table_pdf(src, [SOCIOS, [['Mes', 'Total'], ['Enero', '100']]])
    out = tmp_path / 'out.xlsx'
    tables = convert_pdf_to_excel(src, out)

    assert tables == 2
    wb = load_workbook(out)
    assert len(wb.sheetnames) == 2
    # El nombre de hoja lleva número de tabla y página de origen.
    assert wb.sheetnames[0].startswith('Tabla 1')
    assert '(pag 1)' in wb.sheetnames[0]


def test_pdf_without_tables_rejected_with_clear_message(tmp_path):
    src = tmp_path / 'in.pdf'
    make_text_pdf(src)
    with pytest.raises(ValueError, match='No tables were detected'):
        convert_pdf_to_excel(src, tmp_path / 'out.xlsx')


def test_handler_registers_xlsx_output(tmp_path):
    src = tmp_path / 'in.pdf'
    make_table_pdf(src, [SOCIOS])
    cur = FakeCursor([_original(src, 'cartera.pdf')])
    result = handle_pdf_to_excel({'id': 'j1', 'operation_params': {}}, cur)

    assert result['tables_extracted'] == 1
    assert cur.inserts[0][3] == 'cartera.xlsx'
    assert 'spreadsheetml' in cur.inserts[0][6]


def test_handler_custom_name_strips_xlsx_suffix(tmp_path):
    src = tmp_path / 'in.pdf'
    make_table_pdf(src, [SOCIOS])
    cur = FakeCursor([_original(src)])
    handle_pdf_to_excel({'id': 'j1', 'operation_params': {'output_name': 'datos.xlsx'}}, cur)
    assert cur.inserts[0][3] == 'datos.xlsx'
